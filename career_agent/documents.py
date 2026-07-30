from __future__ import annotations

import io
import re
from pathlib import Path

from .models import ResumeProfile


KNOWN_SKILLS = (
    "Python",
    "LangGraph",
    "LangChain",
    "FastAPI",
    "Pydantic",
    "TypeScript",
    "Node.js",
    "MCP",
    "Docker",
    "SQLite",
    "OpenTelemetry",
    "RAG",
)


def parse_resume_text(text: str) -> ResumeProfile:
    normalized = text.strip()
    if not normalized:
        raise ValueError("resume text is empty")
    heading = re.search(r"^#?\s*([^\n|｜]+)", normalized)
    graduation = re.search(r"(20\d{2})\s*届", normalized)
    location = next((city for city in ("广州", "深圳", "上海", "北京", "杭州", "远程") if city in normalized), "")
    education = next((value for value in ("博士", "硕士", "本科", "大专") if value in normalized), "")
    skills = tuple(skill for skill in KNOWN_SKILLS if re.search(rf"(?<![A-Za-z]){re.escape(skill)}(?![A-Za-z])", normalized, re.I))
    evidence = tuple(
        line.strip(" -*")
        for line in normalized.splitlines()
        if line.strip() and any(skill.casefold() in line.casefold() for skill in skills)
    )
    return ResumeProfile(
        name=(heading.group(1).strip() if heading else "未命名候选人"),
        education=education,
        graduation_year=int(graduation.group(1)) if graduation else 0,
        location=location,
        skills=skills,
        evidence=evidence,
    )


def extract_resume_text(path: str | Path, max_bytes: int = 5 * 1024 * 1024) -> str:
    file_path = Path(path)
    if file_path.stat().st_size > max_bytes:
        raise ValueError("resume file exceeds 5 MiB limit")
    suffix = file_path.suffix.lower()
    if suffix in (".md", ".txt"):
        return file_path.read_text(encoding="utf-8")
    payload = file_path.read_bytes()
    if suffix == ".pdf":
        try:
            import fitz
        except ImportError as error:
            raise RuntimeError("install PyMuPDF to parse PDF resumes") from error
        document = fitz.open(stream=payload, filetype="pdf")
        if document.needs_pass:
            raise ValueError("encrypted PDF resumes are not supported")
        return "\n".join(page.get_text("text") for page in document)
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("install python-docx to parse DOCX resumes") from error
        document = Document(io.BytesIO(payload))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_lines = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + table_lines)
    raise ValueError(f"unsupported resume format: {suffix}")
